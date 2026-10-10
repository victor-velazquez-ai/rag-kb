"""File parsers: turn one file into structural blocks.

Every parser returns a ParsedDocument: a title, the blocks in reading order, the page count when the
format has pages, and the media type. Conventions shared by all formats:

- Headings are emitted as blocks of kind "heading" (so passage text keeps them) and are also carried
  in the ``heading_path`` of every block below them, the heading's own block included.
- A leading level-one heading in Markdown and HTML is the document title: it becomes
  ``ParsedDocument.title`` and is left out of the blocks and heading paths, so sections read
  "Leave > Annual leave" rather than repeating the title.
- Tables become one block of kind "table": one line per row, cells joined by " | ", header row first.
  CSV and Excel produce one table block per group of rows, each repeating the header row.
- Lists become one block of kind "list" with one item per line.

Failures raise UnsupportedFileTypeError, EmptyDocumentError or ParseError, each with a plain message
that is safe to show to the person who uploaded the file.
"""

from __future__ import annotations

import csv
import io
import logging
import re
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import yaml
from bs4 import BeautifulSoup, Comment, NavigableString, Tag
from docx import Document as open_docx
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from markdown_it import MarkdownIt
from markdown_it.token import Token
from openpyxl import load_workbook
from pypdf import PdfReader

from ragkb.ingest.chunker import estimate_tokens
from ragkb.models import Block

log = logging.getLogger(__name__)

# --------------------------------------------------------------------------------------------------
# media types and errors
# --------------------------------------------------------------------------------------------------

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MARKDOWN = "text/markdown"
HTML = "text/html"
TEXT = "text/plain"
CSV = "text/csv"
TSV = "text/tab-separated-values"

_MEDIA_TYPES: dict[str, str] = {
    ".pdf": PDF,
    ".docx": DOCX,
    ".md": MARKDOWN,
    ".markdown": MARKDOWN,
    ".html": HTML,
    ".htm": HTML,
    ".txt": TEXT,
    ".csv": CSV,
    ".tsv": TSV,
    ".xlsx": XLSX,
}

SUPPORTED_SUFFIXES: frozenset[str] = frozenset(_MEDIA_TYPES)

_LEGACY_HINTS = {
    ".doc": "Older Word files (.doc) are not supported. Save the file as .docx and upload it again.",
    ".xls": "Older Excel files (.xls) are not supported. Save the file as .xlsx and upload it again.",
}

_FORMAT_LABELS = {
    PDF: "PDF",
    DOCX: "Word",
    XLSX: "Excel",
    MARKDOWN: "Markdown",
    HTML: "HTML",
    TEXT: "text",
    CSV: "CSV",
    TSV: "TSV",
}

# Rows per CSV or Excel block are grouped up to about this many tokens, so a block usually fits one
# chunk; the chunker still splits a larger block by rows and repeats the header.
ROW_GROUP_TOKENS = 240


class IngestError(Exception):
    """Base class for ingestion errors. ``code`` matches the API error envelope."""

    code = "ingest_error"

    def __init__(self, message: str):
        super().__init__(message)
        self.message = message


class UnsupportedFileTypeError(IngestError):
    code = "unsupported_file_type"


class EmptyDocumentError(IngestError):
    code = "empty_document"


class ParseError(IngestError):
    code = "parse_error"


@dataclass
class ParsedDocument:
    title: str
    blocks: list[Block]
    pages: int | None
    media_type: str


def media_type_for(filename: str) -> str | None:
    """The media type for a file name, from its suffix; None when the type is not supported."""
    return _MEDIA_TYPES.get(Path(filename).suffix.lower())


def unsupported_message(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in _LEGACY_HINTS:
        return _LEGACY_HINTS[suffix]
    shown = suffix or "files without an extension"
    supported = ", ".join(sorted(SUPPORTED_SUFFIXES))
    return f"This file type ({shown}) is not supported. Supported types: {supported}."


def title_from_filename(filename: str) -> str:
    """A readable title from a file name: 'travel-and_expense-policy.md' -> 'Travel and expense policy'."""
    stem = Path(filename).stem
    text = re.sub(r"[_\-.]+", " ", stem)
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return "Untitled"
    if text.islower():
        text = text[0].upper() + text[1:]
    return text


def parse(path: Path, media_type: str | None = None) -> ParsedDocument:
    """Parse one file into blocks. Raises UnsupportedFileTypeError, EmptyDocumentError or ParseError."""
    path = Path(path)
    media_type = media_type or media_type_for(path.name)
    handler = _HANDLERS.get(media_type or "")
    if handler is None:
        raise UnsupportedFileTypeError(unsupported_message(path.name))
    label = _FORMAT_LABELS.get(media_type or "", "")
    try:
        title, blocks, pages = handler(path)
    except IngestError:
        raise
    except Exception as exc:  # any library failure becomes one plain message
        log.warning("parse failed for %s: %s: %s", path.name, type(exc).__name__, exc)
        raise ParseError(
            f"The {label} file could not be read. It may be damaged or saved in an unexpected format."
        ) from exc
    blocks = [b for b in blocks if _has_words(b.text)]
    if not any(b.kind != "heading" for b in blocks):
        raise EmptyDocumentError("The file contains no readable text.")
    title = _clean_title(title) or title_from_filename(path.name)
    return ParsedDocument(title=title, blocks=blocks, pages=pages, media_type=media_type or "")


# --------------------------------------------------------------------------------------------------
# shared helpers
# --------------------------------------------------------------------------------------------------

_WORD = re.compile(r"\w")
_SPACES = re.compile(r"[ \t\f\v\xa0]+")
_LIST_ITEM = re.compile(r"^(?:[-*\u2022\u2023\u25e6\u25aa\u25cf\xb7]|\d{1,3}[.)]|[a-zA-Z][.)])\s+\S")
_NUMBERED_HEADING = re.compile(r"^(\d+(?:\.\d+)*)\.?\s+\S")
_MD_HEADING_LINE = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")


def _has_words(text: str) -> bool:
    return bool(_WORD.search(text))


def _collapse(text: str) -> str:
    """Collapse runs of spaces inside lines, trim lines, drop empty lines."""
    lines = (_SPACES.sub(" ", line).strip() for line in text.splitlines())
    return "\n".join(line for line in lines if line)


def _one_line(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _clean_title(title: str | None) -> str:
    title = _one_line(title or "")
    return title[:200]


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        text = raw.decode("utf-16")
    else:
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            text = raw.decode("latin-1")
    return text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")


def _looks_like_heading(line: str) -> bool:
    """A conservative guess for formats without structure (PDF, plain text): a short line with no
    closing punctuation that starts with a capital letter or a section number."""
    line = line.strip()
    if not line or len(line) > 70 or len(line.split()) > 8:
        return False
    if line[-1] in ".,;:!?" or (_LIST_ITEM.match(line) and not _NUMBERED_HEADING.match(line)):
        return False
    if not re.search(r"[A-Za-z]", line):
        return False
    return line[0].isupper() or line[0].isdigit()


def _heading_level(line: str) -> int:
    m = _NUMBERED_HEADING.match(line)
    if m:
        return min(6, m.group(1).count(".") + 1)
    return 1


class _Sink:
    """Collects blocks and tracks the heading path."""

    def __init__(self, title_from_first_h1: bool):
        self.blocks: list[Block] = []
        self.title = ""
        self._stack: list[tuple[int, str]] = []
        self._title_open = title_from_first_h1

    @property
    def path(self) -> list[str]:
        return [text for _, text in self._stack]

    def heading(self, level: int, text: str, page: int | None = None) -> None:
        text = _one_line(text)
        if not _has_words(text):
            return
        if self._title_open and level == 1 and not self.blocks:
            self.title = text
            self._title_open = False
            return
        self._title_open = False
        while self._stack and self._stack[-1][0] >= level:
            self._stack.pop()
        self._stack.append((level, text))
        self.blocks.append(Block(text=text, heading_path=self.path, page=page, kind="heading"))

    def add(self, text: str, kind: str = "paragraph", page: int | None = None) -> None:
        text = text.strip("\n").rstrip() if kind == "code" else text.strip()
        if not _has_words(text):
            return
        self._title_open = False
        self.blocks.append(Block(text=text, heading_path=self.path, page=page, kind=kind))  # type: ignore[arg-type]


# --------------------------------------------------------------------------------------------------
# Markdown
# --------------------------------------------------------------------------------------------------

_MD = MarkdownIt("commonmark").enable("table")
_FRONT_MATTER = re.compile(r"\A---[ \t]*\n(.*?)\n(?:---|\.\.\.)[ \t]*(?:\n|\Z)", re.DOTALL)


def _parse_markdown(path: Path) -> tuple[str, list[Block], int | None]:
    text = _read_text(path)
    front_title = ""
    m = _FRONT_MATTER.match(text)
    if m:
        try:
            meta = yaml.safe_load(m.group(1))
        except yaml.YAMLError:
            meta = None
        if isinstance(meta, dict) and isinstance(meta.get("title"), str):
            front_title = meta["title"]
        text = text[m.end() :]
    tokens = _MD.parse(text)
    sink = _Sink(title_from_first_h1=True)
    i = 0
    while i < len(tokens):
        tok = tokens[i]
        if tok.type == "heading_open":
            sink.heading(int(tok.tag[1]), _inline_text(tokens[i + 1]))
            i += 3
        elif tok.type == "paragraph_open":
            sink.add(_inline_text(tokens[i + 1]))
            i += 3
        elif tok.type in ("bullet_list_open", "ordered_list_open"):
            end = _matching_close(tokens, i)
            sink.add(_md_list(tokens[i : end + 1]), kind="list")
            i = end + 1
        elif tok.type == "table_open":
            end = _matching_close(tokens, i)
            sink.add(_md_table(tokens[i : end + 1]), kind="table")
            i = end + 1
        elif tok.type in ("fence", "code_block"):
            sink.add(tok.content, kind="code")
            i += 1
        elif tok.type == "html_block":
            sink.add(_html_fragment_text(tok.content))
            i += 1
        else:  # blockquote and hr wrappers: their inner blocks are handled on the next tokens
            i += 1
    return front_title or sink.title, sink.blocks, None


def _matching_close(tokens: list[Token], start: int) -> int:
    depth = 0
    for j in range(start, len(tokens)):
        depth += tokens[j].nesting
        if depth == 0:
            return j
    return len(tokens) - 1


def _inline_text(tok: Token) -> str:
    parts: list[str] = []
    for child in tok.children or []:
        if child.type in ("text", "code_inline"):
            parts.append(child.content)
        elif child.type == "softbreak":
            parts.append(" ")
        elif child.type == "hardbreak":
            parts.append("\n")
        elif child.type == "image":
            parts.append(child.content or "")
    return _collapse("".join(parts))


def _md_list(tokens: list[Token]) -> str:
    lines: list[str] = []
    kinds: list[str] = []
    counters: list[int] = []
    prefix: str | None = None
    for tok in tokens:
        if tok.type == "bullet_list_open":
            kinds.append("ul")
            counters.append(0)
        elif tok.type == "ordered_list_open":
            kinds.append("ol")
            start = tok.attrGet("start")
            counters.append(int(start) - 1 if start is not None else 0)
        elif tok.type in ("bullet_list_close", "ordered_list_close"):
            kinds.pop()
            counters.pop()
        elif tok.type == "list_item_open":
            indent = "  " * (len(kinds) - 1)
            if kinds[-1] == "ol":
                counters[-1] += 1
                prefix = f"{indent}{counters[-1]}. "
            else:
                prefix = f"{indent}- "
        elif tok.type == "inline":
            text = _inline_text(tok).replace("\n", " ")
            if prefix is not None:
                lines.append(prefix + text)
                prefix = None
            else:
                lines.append("  " * len(kinds) + text)
        elif tok.type in ("fence", "code_block"):
            indent = "  " * len(kinds)
            lines.extend(indent + line for line in tok.content.rstrip("\n").splitlines())
    return "\n".join(line for line in lines if line.strip())


def _md_table(tokens: list[Token]) -> str:
    rows: list[str] = []
    cells: list[str] = []
    for tok in tokens:
        if tok.type == "tr_open":
            cells = []
        elif tok.type == "inline":
            cells.append(_inline_text(tok).replace("\n", " "))
        elif tok.type == "tr_close":
            row = _format_row(cells)
            if row:
                rows.append(row)
    return "\n".join(rows)


def _format_row(cells: list[str]) -> str:
    cells = [_one_line(c) for c in cells]
    while cells and not cells[-1]:
        cells.pop()
    return " | ".join(cells) if any(cells) else ""


# --------------------------------------------------------------------------------------------------
# HTML
# --------------------------------------------------------------------------------------------------

_HTML_SKIP = {
    "script",
    "style",
    "nav",
    "noscript",
    "template",
    "svg",
    "iframe",
    "form",
    "button",
    "select",
    "footer",
    "canvas",
    "object",
    "embed",
    "head",
    "audio",
    "video",
    "map",
}
_HTML_BLOCKISH = {
    "div",
    "section",
    "article",
    "main",
    "body",
    "html",
    "header",
    "aside",
    "figure",
    "figcaption",
    "details",
    "summary",
    "blockquote",
    "dl",
    "dt",
    "dd",
    "li",
    "fieldset",
    "center",
    "address",
    "hgroup",
}
_BR_MARK = "\ue000"


def _parse_html(path: Path) -> tuple[str, list[Block], int | None]:
    soup = BeautifulSoup(_read_text(path), "html.parser")
    page_title = soup.title.get_text() if soup.title else ""
    for br in soup.find_all("br"):
        br.replace_with(_BR_MARK)
    root = soup.find("main") or soup.body or soup
    walker = _HtmlWalker()
    walker.walk(root)
    walker.flush()
    return walker.sink.title or page_title, walker.sink.blocks, None


def _html_text(node: Tag | str) -> str:
    raw = node.get_text() if isinstance(node, Tag) else str(node)
    raw = re.sub(r"\s+", " ", raw)
    return _collapse(raw.replace(_BR_MARK, "\n"))


def _html_fragment_text(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup.find_all(list(_HTML_SKIP)):
        tag.decompose()
    return _one_line(soup.get_text(" "))


class _HtmlWalker:
    def __init__(self) -> None:
        self.sink = _Sink(title_from_first_h1=True)
        self.buffer: list[str] = []

    def flush(self) -> None:
        if self.buffer:
            self.sink.add(_html_text("".join(self.buffer)))
            self.buffer = []

    def walk(self, node: Tag) -> None:
        for child in node.children:
            if isinstance(child, Comment):
                continue
            if isinstance(child, NavigableString):
                self.buffer.append(str(child))
                continue
            if not isinstance(child, Tag):
                continue
            name = (child.name or "").lower()
            if name in _HTML_SKIP:
                continue
            if re.fullmatch(r"h[1-6]", name):
                self.flush()
                self.sink.heading(int(name[1]), _html_text(child).replace("\n", " "))
            elif name == "p":
                self.flush()
                self.sink.add(_html_text(child))
            elif name in ("ul", "ol"):
                self.flush()
                self.sink.add(_html_list(child), kind="list")
            elif name == "table":
                self.flush()
                self.sink.add(_html_table(child), kind="table")
            elif name == "pre":
                self.flush()
                self.sink.add(child.get_text().replace(_BR_MARK, "\n"), kind="code")
            elif name in _HTML_BLOCKISH or child.find(_block_tag_names()) is not None:
                self.flush()
                self.walk(child)
                self.flush()
            else:
                self.buffer.append(child.get_text())


def _block_tag_names() -> list[str]:
    return ["p", "ul", "ol", "table", "pre", "h1", "h2", "h3", "h4", "h5", "h6", "div", "section"]


def _html_list(lst: Tag, depth: int = 0) -> str:
    ordered = (lst.name or "").lower() == "ol"
    try:
        number = int(lst.get("start", 1))
    except (TypeError, ValueError):
        number = 1
    lines: list[str] = []
    for li in lst.find_all("li", recursive=False):
        nested = [n for n in li.find_all(["ul", "ol"]) if n.find_parent("li") is li]
        for n in nested:
            n.extract()
        text = _html_text(li).replace("\n", " ")
        marker = f"{number}." if ordered else "-"
        number += 1
        if text:
            lines.append(f"{'  ' * depth}{marker} {text}")
        for n in nested:
            inner = _html_list(n, depth + 1)
            if inner:
                lines.append(inner)
    return "\n".join(lines)


def _html_table(table: Tag) -> str:
    rows: list[str] = []
    for tr in table.find_all("tr"):
        if tr.find_parent("table") is not table:
            continue
        cells = [_html_text(c).replace("\n", " ") for c in tr.find_all(["th", "td"], recursive=False)]
        row = _format_row(cells)
        if row:
            rows.append(row)
    return "\n".join(rows)


# --------------------------------------------------------------------------------------------------
# Word
# --------------------------------------------------------------------------------------------------

_DOCX_HEADING = re.compile(r"heading\s*(\d)", re.IGNORECASE)


def _parse_docx(path: Path) -> tuple[str, list[Block], int | None]:
    document = open_docx(str(path))
    sink = _Sink(title_from_first_h1=False)
    title = ""
    list_lines: list[str] = []
    numbered = 0

    def flush_list() -> None:
        nonlocal numbered
        if list_lines:
            sink.add("\n".join(list_lines), kind="list")
            list_lines.clear()
        numbered = 0

    for item in document.iter_inner_content():
        if isinstance(item, DocxTable):
            flush_list()
            sink.add(_docx_table(item), kind="table")
            continue
        if not isinstance(item, DocxParagraph):
            continue
        text = _collapse(item.text)
        if not text:
            continue
        style = (item.style.name if item.style is not None else "") or ""
        if style.lower() == "title":
            flush_list()
            if not title and not sink.blocks:
                title = _one_line(text)
            else:
                sink.heading(1, text)
            continue
        heading = _DOCX_HEADING.match(style)
        if heading:
            flush_list()
            sink.heading(int(heading.group(1)), text)
            continue
        if _docx_is_list(item, style):
            if "number" in style.lower():
                numbered += 1
                list_lines.append(f"{numbered}. {_one_line(text)}")
            else:
                list_lines.append(f"- {_one_line(text)}")
            continue
        flush_list()
        sink.add(text)
    flush_list()
    title = title or (document.core_properties.title or "")
    return title, sink.blocks, None


def _docx_is_list(paragraph: DocxParagraph, style: str) -> bool:
    if "list" in style.lower():
        return True
    p_pr = paragraph._p.pPr  # python-docx exposes numbering only through the XML
    return p_pr is not None and p_pr.numPr is not None


def _docx_table(table: DocxTable) -> str:
    rows: list[str] = []
    for row in table.rows:
        seen: set[int] = set()
        cells: list[str] = []
        for cell in row.cells:
            key = id(cell._tc)  # merged cells repeat the same underlying element
            if key in seen:
                continue
            seen.add(key)
            cells.append(_one_line(cell.text))
        line = _format_row(cells)
        if line:
            rows.append(line)
    return "\n".join(rows)


# --------------------------------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------------------------------

_PAGE_NUMBER = re.compile(r"^(?:page\s*)?\d{1,4}(?:\s*(?:of|/)\s*\d{1,4})?$|^-\s*\d{1,4}\s*-$", re.IGNORECASE)
_BAD_PDF_TITLES = re.compile(
    r"^(?:untitled|microsoft word\b|document\d*$)|\.(?:docx?|pdf|pptx?|xlsx?|txt)$", re.IGNORECASE
)
_SCANNED_MESSAGE = (
    "This PDF has no extractable text. It looks like a scanned document, and scanned PDFs need OCR, "
    "which is not supported yet."
)


def _parse_pdf(path: Path) -> tuple[str, list[Block], int | None]:
    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise ParseError("The PDF could not be opened. It may be damaged.") from exc
    if reader.is_encrypted:
        try:
            unlocked = bool(reader.decrypt(""))
        except Exception:  # noqa: BLE001 - any decryption failure means locked
            unlocked = False
        if not unlocked:
            raise ParseError("The PDF is password protected. Remove the password and upload it again.")
    raw_pages: list[str] = []
    for page in reader.pages:
        try:
            raw_pages.append(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001 - one bad page should not lose the document
            log.warning("pdf page extraction failed in %s: %s", path.name, exc)
            raw_pages.append("")
    page_count = len(raw_pages)
    page_lines = _drop_running_lines([_pdf_lines(t) for t in raw_pages])
    chars = sum(len("".join(line.split())) for lines in page_lines for line in lines)
    if chars < 20 or chars < 10 * max(page_count, 1):
        raise EmptyDocumentError(_SCANNED_MESSAGE)

    items = [
        (number, kind, text)
        for number, lines in enumerate(page_lines, start=1)
        for kind, text in _pdf_paragraphs(lines)
    ]
    try:
        meta_title = _one_line((reader.metadata.title if reader.metadata else None) or "")
    except Exception:  # noqa: BLE001 - unreadable metadata only costs the title
        meta_title = ""
    title = meta_title if meta_title and not _BAD_PDF_TITLES.search(meta_title) else ""
    # a leading heading is the document title when the metadata has none or says the same
    if items and items[0][1] == "heading" and (not title or items[0][2].lower() == title.lower()):
        title = title or items[0][2]
        items = items[1:]

    sink = _Sink(title_from_first_h1=False)
    for number, kind, text in items:
        if kind == "heading":
            sink.heading(_heading_level(text), text, page=number)
        else:
            sink.add(text, kind=kind, page=number)
    return title, sink.blocks, page_count


def _pdf_lines(text: str) -> list[str]:
    lines: list[str] = []
    for raw in text.replace("\r", "\n").split("\n"):
        line = _SPACES.sub(" ", raw).strip()
        if _PAGE_NUMBER.match(line):
            continue
        lines.append(line)
    # join words broken across lines: "manage-" + "ment" -> "management"
    joined: list[str] = []
    for line in lines:
        if (
            joined
            and line
            and line[0].islower()
            and re.search(r"[A-Za-z]-$", joined[-1])
            and not re.search(r"\s-$", joined[-1])
        ):
            joined[-1] = joined[-1][:-1] + line
        else:
            joined.append(line)
    return joined


def _drop_running_lines(pages: list[list[str]]) -> list[list[str]]:
    """Remove running headers and footers: lines at the top or bottom of most pages."""
    if len(pages) < 3:
        return pages

    def edges(lines: list[str]) -> set[str]:
        content = [line for line in lines if line]
        return {re.sub(r"\d+", "#", line) for line in content[:2] + content[-2:]}

    counts: Counter[str] = Counter()
    for lines in pages:
        counts.update(edges(lines))
    running = {key for key, n in counts.items() if n >= max(3, int(0.6 * len(pages)))}
    if not running:
        return pages
    cleaned: list[list[str]] = []
    for lines in pages:
        content = [line for line in lines if line]
        edge_lines = set(content[:2] + content[-2:])
        cleaned.append(
            [line for line in lines if not (line in edge_lines and re.sub(r"\d+", "#", line) in running)]
        )
    return cleaned


def _pdf_paragraphs(lines: list[str]) -> list[tuple[str, str]]:
    """Group the lines of one page into (kind, text) paragraphs, list blocks and headings."""
    lengths = sorted(len(line) for line in lines if line)
    width = lengths[int(0.8 * (len(lengths) - 1))] if lengths else 0
    groups: list[list[str]] = []
    current: list[str] = []
    for line in lines:
        if not line:
            if current:
                groups.append(current)
                current = []
            continue
        if current:
            prev = current[-1]
            # In wrapped text only the last line of a paragraph is short, so a short line ends one.
            short = len(prev) < 0.6 * width
            ends_sentence = prev[-1] in ".!?:" and len(prev) < 0.85 * width
            if short or ends_sentence or _LIST_ITEM.match(line):
                groups.append(current)
                current = []
        current.append(line)
    if current:
        groups.append(current)

    # consecutive list items become one list block, one item per line
    out: list[tuple[str, str]] = []
    for index, group in enumerate(groups):
        if _LIST_ITEM.match(group[0]):
            item = _join_lines(group)
            if out and out[-1][0] == "list":
                out[-1] = ("list", out[-1][1] + "\n" + item)
            else:
                out.append(("list", item))
            continue
        text = _join_lines(group)
        next_exists = index + 1 < len(groups)
        if len(group) == 1 and next_exists and _looks_like_heading(text):
            out.append(("heading", text))
        else:
            out.append(("paragraph", text))
    return out


def _join_lines(lines: list[str]) -> str:
    text = " ".join(line.strip() for line in lines)
    return _SPACES.sub(" ", text).strip()


# --------------------------------------------------------------------------------------------------
# plain text
# --------------------------------------------------------------------------------------------------


def _parse_text(path: Path) -> tuple[str, list[Block], int | None]:
    text = _read_text(path)
    paragraphs = [p for p in re.split(r"\n[ \t]*\n", text) if p.strip()]
    title = ""
    if len(paragraphs) > 1:
        first = paragraphs[0].strip()
        if "\n" not in first and len(first) <= 80 and first[-1] not in ".!?,;:" and _has_words(first):
            title = first.lstrip("#").strip()
            paragraphs = paragraphs[1:]
    sink = _Sink(title_from_first_h1=False)
    for index, para in enumerate(paragraphs):
        lines = [line.strip() for line in para.splitlines() if line.strip()]
        if not lines:
            continue
        md_heading = _MD_HEADING_LINE.match(lines[0])
        if md_heading:
            sink.heading(len(md_heading.group(1)), md_heading.group(2))
            lines = lines[1:]
            if not lines:
                continue
        elif len(lines) == 1 and index + 1 < len(paragraphs) and _looks_like_heading(lines[0]):
            sink.heading(_heading_level(lines[0]), lines[0])
            continue
        # lines before the first list item are a lead-in paragraph; the rest is the list
        first_item = next((i for i, line in enumerate(lines) if _LIST_ITEM.match(line)), len(lines))
        if first_item:
            sink.add(" ".join(lines[:first_item]))
        items: list[str] = []
        for line in lines[first_item:]:
            if _LIST_ITEM.match(line) or not items:
                items.append(line)
            else:
                items[-1] += " " + line
        if items:
            sink.add("\n".join(items), kind="list")
    return title, sink.blocks, None


# --------------------------------------------------------------------------------------------------
# CSV and Excel
# --------------------------------------------------------------------------------------------------


def _parse_csv(path: Path) -> tuple[str, list[Block], int | None]:
    return _parse_delimited(path, delimiter=None)


def _parse_tsv(path: Path) -> tuple[str, list[Block], int | None]:
    return _parse_delimited(path, delimiter="\t")


def _parse_delimited(path: Path, delimiter: str | None) -> tuple[str, list[Block], int | None]:
    text = _read_text(path)
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows = [[_one_line(cell) for cell in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    return "", _row_group_blocks(rows, heading_path=[]), None


def _parse_xlsx(path: Path) -> tuple[str, list[Block], int | None]:
    workbook = load_workbook(str(path), read_only=True, data_only=True)
    try:
        sheets: list[tuple[str, list[list[str]]]] = []
        for sheet in workbook.worksheets:
            rows = [[_format_cell(v) for v in values] for values in sheet.iter_rows(values_only=True)]
            if any(any(row) for row in rows):
                sheets.append((sheet.title, rows))
        title = workbook.properties.title or ""
    finally:
        workbook.close()
    blocks: list[Block] = []
    for name, rows in sheets:
        default_name = re.fullmatch(r"sheet\s*\d*", name.strip(), re.IGNORECASE) is not None
        path_ = [] if default_name and len(sheets) == 1 else [name]
        blocks.extend(_row_group_blocks(rows, heading_path=path_))
    return title, blocks, None


def _format_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:.10g}"
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat(timespec="minutes")
    return _one_line(str(value))


def _row_group_blocks(rows: list[list[str]], heading_path: list[str]) -> list[Block]:
    rows = [row for row in rows if any(row)]
    if not rows:
        return []
    width = max(max((i + 1 for i, cell in enumerate(row) if cell), default=0) for row in rows)
    rows = [row[:width] for row in rows]
    header = _format_row(rows[0])
    body = [line for line in (_format_row(row) for row in rows[1:]) if line]
    if not body:
        return [Block(text=header, heading_path=heading_path, kind="table")]
    header_tokens = estimate_tokens(header)
    blocks: list[Block] = []
    group: list[str] = []
    tokens = header_tokens
    for line in body:
        line_tokens = estimate_tokens(line)
        if group and tokens + line_tokens > ROW_GROUP_TOKENS:
            blocks.append(Block(text="\n".join([header, *group]), heading_path=heading_path, kind="table"))
            group, tokens = [], header_tokens
        group.append(line)
        tokens += line_tokens
    if group:
        blocks.append(Block(text="\n".join([header, *group]), heading_path=heading_path, kind="table"))
    return blocks


_HANDLERS = {
    PDF: _parse_pdf,
    DOCX: _parse_docx,
    MARKDOWN: _parse_markdown,
    HTML: _parse_html,
    TEXT: _parse_text,
    CSV: _parse_csv,
    TSV: _parse_tsv,
    XLSX: _parse_xlsx,
}
