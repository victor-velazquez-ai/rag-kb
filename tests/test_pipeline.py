"""Ingestion pipeline: identity by (source path, content hash), durability on failure, failure rules,
progress, batching, cancellation, concurrency, page citations, versions."""

from __future__ import annotations

import hashlib
import re
import threading
import time
from collections.abc import Sequence
from pathlib import Path
from typing import Any

import pytest
from docx import Document as WordDocument
from openpyxl import Workbook

from ragkb.config import Settings
from ragkb.ingest.parsers import (
    EmptyDocumentError,
    EncryptedFileError,
    FileTooLargeError,
    UnsupportedFileTypeError,
)
from ragkb.ingest.pipeline import (
    IngestCancelled,
    IngestPipeline,
    IngestResult,
    NotFoundError,
    normalize_source_path,
    safe_filename,
)
from ragkb.models import Collection, DocumentStatus
from ragkb.retrieval.embedder import EmbedderError, FakeEmbedder
from ragkb.store.lancedb_store import LanceDBStore
from ragkb.store.sqlite_meta import SQLiteMeta
from tests.test_parsers import encrypted_pdf

VERSION = re.compile(r"v\d{8}\.\d+")


def markdown(title: str, sections: dict[str, str]) -> bytes:
    parts = [f"# {title}", ""]
    for heading, body in sections.items():
        parts += [f"## {heading}", "", body, ""]
    return "\n".join(parts).encode("utf-8")


def filler(topic: str, n: int = 12) -> str:
    return " ".join(
        f"Rule {i} about {topic} applies to every team and every office location." for i in range(n)
    )


V1 = markdown(
    "Travel Rules",
    {
        "Flights": "Economy class is required for flights under six hours. " + filler("flights"),
        "Hotels": "Hotels in capital cities are capped at 210 euros per night. " + filler("hotels"),
        "Meals": "The zebrafruit allowance is 70 euros per day abroad. " + filler("meals"),
        "Trains": "Trains are preferred for journeys under 500 kilometres. " + filler("trains"),
        "Cars": "Private cars are reimbursed at 30 cents per kilometre. " + filler("cars"),
    },
)
V2 = markdown(
    "Travel Rules", {"Flights": "Premium economy is allowed for flights over six hours. " + filler("flights")}
)
SYLLABLES = ["ka", "lo", "mi", "nu", "pe", "ri", "so", "tu"]
CODE_WORDS = [f"{SYLLABLES[i // 8]}{SYLLABLES[i % 8]}zen" for i in range(40)]


class FlakyEmbedder(FakeEmbedder):
    def __init__(self) -> None:
        super().__init__(dims=384)
        self.fail = False

    def embed(self, texts: Sequence[str]) -> list[list[float]]:
        if self.fail:
            raise EmbedderError("The embedding model fake failed: simulated outage")
        return super().embed(texts)


class GatedEmbedder(FakeEmbedder):
    """While the gate is closed, ``embed`` waits on batches that contain ``block_on`` (any batch when it
    is empty), so a test can act while one upload is embedding and others go through."""

    def __init__(self) -> None:
        super().__init__(dims=384)
        self.gate = threading.Event()
        self.gate.set()
        self.block_on = ""
        self.waiting = threading.Event()
        self.batches: list[int] = []

    def embed(self, texts: Sequence[str]) -> list[list[float]]:
        self.batches.append(len(texts))
        if not self.gate.is_set() and any(self.block_on in t for t in texts):
            self.waiting.set()
            assert self.gate.wait(20), "the test never opened the gate"
        return super().embed(texts)


@pytest.fixture
def flaky() -> FlakyEmbedder:
    return FlakyEmbedder()


@pytest.fixture
def flaky_pipeline(
    settings: Settings, meta: SQLiteMeta, store: LanceDBStore, flaky: FlakyEmbedder
) -> IngestPipeline:
    return IngestPipeline(settings, meta, store, flaky)


@pytest.fixture
def gated() -> GatedEmbedder:
    return GatedEmbedder()


@pytest.fixture
def gated_pipeline(
    settings: Settings, meta: SQLiteMeta, store: LanceDBStore, gated: GatedEmbedder
) -> IngestPipeline:
    return IngestPipeline(settings, meta, store, gated)


def texts(store: LanceDBStore, doc_id: str) -> str:
    return "\n".join(c.text for c in store.chunks_for(doc_id))


def in_thread(target: Any, *args: Any, **kwargs: Any) -> tuple[threading.Thread, dict[str, Any]]:
    """Run ``target`` in a thread; the dict receives its "result" or "error"."""
    out: dict[str, Any] = {}

    def run() -> None:
        try:
            out["result"] = target(*args, **kwargs)
        except BaseException as exc:  # noqa: BLE001 - report any failure from the thread
            out["error"] = exc

    thread = threading.Thread(target=run)
    thread.start()
    return thread, out


def fact_pdf(pdf_builder, pages: int) -> bytes:
    """One distinct fact per page, with a little filler, so each page is short."""
    return pdf_builder(
        [
            [
                f"The vault {n} code word is {CODE_WORDS[n - 1]}.",
                "",
                "Visitors sign the register at the front desk before they enter the archive rooms.",
            ]
            for n in range(1, pages + 1)
        ]
    )


# ------------------------------------------------------------------------------------------- ingest


def test_ingest_bytes_indexes_a_new_document(
    pipeline: IngestPipeline,
    meta: SQLiteMeta,
    store: LanceDBStore,
    settings: Settings,
    collection: Collection,
) -> None:
    result = pipeline.ingest_bytes(collection.id, V1, "travel.md")
    assert isinstance(result, IngestResult)
    assert (result.already_indexed, result.replaced_doc_id, result.replaced_title) == (False, None, None)
    doc = result.document
    assert doc.status == DocumentStatus.ready and doc.error is None
    assert doc.title == "Travel Rules" and doc.filename == "travel.md" and doc.media_type == "text/markdown"
    assert doc.source_path == "travel.md"
    assert doc.size_bytes == len(V1) and doc.content_hash == hashlib.sha256(V1).hexdigest()
    assert doc.pages is None and doc.indexed_at is not None
    assert VERSION.fullmatch(doc.knowledge_version or "")
    chunks = store.chunks_for(doc.id)
    assert doc.chunk_count == len(chunks) >= 5
    assert [c.id for c in chunks] == [f"{doc.id}#{i}" for i in range(len(chunks))]
    assert {c.knowledge_version for c in chunks} == {doc.knowledge_version}
    assert {c.title for c in chunks} == {"Travel Rules"} and {c.collection_id for c in chunks} == {
        collection.id
    }
    assert chunks[1].section == "Hotels" and chunks[1].embed_text.startswith("Travel Rules > Hotels: ")
    assert all((c.page, c.page_start, c.page_end) == (None, None, None) for c in chunks)
    assert all(c.token_estimate == FakeEmbedder().count_tokens(c.text) for c in chunks), "the model's own count"
    kept = settings.uploads_dir / collection.id / doc.id / "travel.md"
    assert kept.read_bytes() == V1 and pipeline.original_path(doc) == kept
    staging = settings.uploads_dir / collection.id / ".staging"
    assert not staging.exists() or not any(staging.iterdir())
    assert meta.get_document(doc.id) == doc
    audit = meta.list_audit(event="ingest")[0]
    assert audit["outcome"] == "indexed" and audit["doc_id"] == doc.id and audit["chunks"] == len(chunks)
    assert {"parse_ms", "chunk_ms", "embed_ms", "write_ms"} <= set(audit), "stage timings are audited"


def test_same_bytes_are_already_indexed(
    pipeline: IngestPipeline, meta: SQLiteMeta, collection: Collection
) -> None:
    first = pipeline.ingest_bytes(collection.id, V1, "travel.md").document
    version = meta.current_version()
    again = pipeline.ingest_bytes(collection.id, V1, "renamed-copy.md", title="Another title")
    assert again.already_indexed and again.document == first
    assert meta.current_version() == version, "no work, no version bump"
    assert len(meta.list_documents(collection.id)) == 1
    assert meta.list_audit(event="ingest")[0]["outcome"] == "duplicate"
    assert pipeline.find_duplicate(collection.id, V1) == first
    assert pipeline.find_duplicate(collection.id, V2) is None


def test_same_bytes_in_another_collection_is_a_new_document(
    pipeline: IngestPipeline, meta: SQLiteMeta, collection: Collection
) -> None:
    other = meta.create_collection("Other")
    first = pipeline.ingest_bytes(collection.id, V1, "travel.md").document
    second = pipeline.ingest_bytes(other.id, V1, "travel.md").document
    assert second.id != first.id and second.collection_id == other.id


def test_same_path_with_new_bytes_replaces_in_place(
    pipeline: IngestPipeline, meta: SQLiteMeta, store: LanceDBStore, collection: Collection
) -> None:
    v1 = pipeline.ingest_bytes(collection.id, V1, "README.md", source_path="docs/travel/README.md").document
    renamed = V2.replace(b"# Travel Rules", b"# Travel Rules 2027")
    result = pipeline.ingest_bytes(collection.id, renamed, "README.md", source_path="docs/travel/README.md")
    v2 = result.document
    assert (result.already_indexed, result.replaced_doc_id, result.replaced_title) == (False, v1.id, "Travel Rules")
    assert v2.id == v1.id and v2.status == DocumentStatus.ready and v2.title == "Travel Rules 2027"
    assert v2.chunk_count == len(store.chunks_for(v2.id)) < v1.chunk_count
    assert v2.content_hash == hashlib.sha256(renamed).hexdigest() and v2.size_bytes == len(renamed)
    assert v2.knowledge_version != v1.knowledge_version
    assert "Premium economy" in texts(store, v2.id) and "zebrafruit" not in texts(store, v2.id)
    assert store.text_search(collection.id, "zebrafruit", 5) == []
    assert [h.chunk.doc_id for h in store.text_search(collection.id, "premium", 5)] == [v2.id]
    assert len(meta.list_documents(collection.id)) == 1
    assert meta.list_audit(event="ingest")[0]["outcome"] == "replaced"


def test_same_title_at_different_paths_keeps_every_file(
    pipeline: IngestPipeline, meta: SQLiteMeta, store: LanceDBStore, collection: Collection
) -> None:
    """Folders of README.md files, or pages that all start with "# Introduction", no longer overwrite
    each other: titles are display only."""
    files = {
        "guides/install/README.md": markdown("Introduction", {"Setup": "Install the agent with the wizard."}),
        "guides/backup/README.md": markdown("Introduction", {"Backups": "Backups run nightly at two."}),
        "handbook/intro.md": markdown("Introduction", {"Welcome": "Welcome to the handbook for staff."}),
        "faq/README.md": markdown("FAQ", {"Billing": "Invoices are sent on the first working day."}),
    }
    results = [
        pipeline.ingest_bytes(collection.id, data, Path(path).name, source_path=path) for path, data in files.items()
    ]
    assert all(r.document.status == DocumentStatus.ready and r.replaced_doc_id is None for r in results)
    docs = meta.list_documents(collection.id)
    assert len(docs) == 4 and len({d.id for d in docs}) == 4
    assert sorted(d.source_path for d in docs) == sorted(files)
    assert [d.title for d in docs].count("Introduction") == 3
    for phrase in ("wizard", "nightly", "handbook", "Invoices"):
        assert store.text_search(collection.id, phrase, 3), phrase


def test_an_explicit_title_is_display_only(pipeline: IngestPipeline, collection: Collection) -> None:
    first = pipeline.ingest_bytes(collection.id, V1, "a.md", title="Policy").document
    second = pipeline.ingest_bytes(collection.id, V2, "b.md", title="Policy").document
    assert first.id != second.id and first.title == second.title == "Policy"
    assert first.metadata["title_source"] == "explicit"


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("docs/guides/README.md", "docs/guides/README.md"),
        ("docs\\guides\\README.md", "docs/guides/README.md"),
        ("./docs//guides/../README.md", "docs/guides/README.md"),
        ("/abs/path.md", "abs/path.md"),
        ("C:\\Users\\me\\notes.txt", "Users/me/notes.txt"),
        ("", ""),
        (None, ""),
    ],
)
def test_normalize_source_path(raw: str | None, expected: str) -> None:
    assert normalize_source_path(raw) == expected


def test_source_path_spellings_name_the_same_file(pipeline: IngestPipeline, collection: Collection) -> None:
    first = pipeline.ingest_bytes(collection.id, V1, "README.md", source_path="docs\\README.md").document
    result = pipeline.ingest_bytes(collection.id, V2, "README.md", source_path="docs/README.md")
    assert result.document.id == first.id and result.replaced_doc_id == first.id


# --------------------------------------------------------------------------------------- durability


def test_failed_replacement_keeps_the_previous_version(
    flaky_pipeline: IngestPipeline,
    flaky: FlakyEmbedder,
    meta: SQLiteMeta,
    store: LanceDBStore,
    settings: Settings,
    collection: Collection,
) -> None:
    v1 = flaky_pipeline.ingest_bytes(collection.id, V1, "travel.md").document
    version = meta.current_version()
    flaky.fail = True
    result = flaky_pipeline.ingest_bytes(collection.id, V2, "travel.md")
    failed = result.document
    assert result.replaced_doc_id is None, "nothing was replaced"
    assert failed.id == v1.id and failed.status == DocumentStatus.failed
    assert "simulated outage" in failed.error and "previous version is still searchable" in failed.error
    assert (failed.chunk_count, failed.content_hash, failed.size_bytes) == (v1.chunk_count, v1.content_hash, len(V1))
    assert meta.current_version() == version
    assert len(store.chunks_for(v1.id)) == v1.chunk_count
    assert [h.chunk.doc_id for h in store.text_search(collection.id, "zebrafruit", 5)] == [v1.id]
    assert (settings.uploads_dir / collection.id / v1.id / "travel.md").read_bytes() == V1
    assert meta.list_audit(event="ingest")[0]["outcome"] == "failed"

    flaky.fail = False
    fixed = flaky_pipeline.ingest_bytes(collection.id, V2, "travel.md")
    assert fixed.document.id == v1.id and fixed.document.status == DocumentStatus.ready
    assert fixed.replaced_doc_id == v1.id and fixed.document.error is None
    assert "Premium economy" in texts(store, v1.id) and "zebrafruit" not in texts(store, v1.id)
    assert (settings.uploads_dir / collection.id / v1.id / "travel.md").read_bytes() == V2


def test_unreadable_replacement_keeps_the_previous_version(
    pipeline: IngestPipeline, store: LanceDBStore, collection: Collection, pdf_builder
) -> None:
    v1 = pipeline.ingest_bytes(collection.id, fact_pdf(pdf_builder, 2), "vaults.pdf").document
    failed = pipeline.ingest_bytes(collection.id, pdf_builder([[], []]), "vaults.pdf").document
    assert failed.id == v1.id and failed.status == DocumentStatus.failed
    assert "scanned" in failed.error and "previous version is still searchable" in failed.error
    assert len(store.chunks_for(v1.id)) == v1.chunk_count
    assert store.text_search(collection.id, CODE_WORDS[0], 5)


def test_new_document_failure_is_recorded(
    pipeline: IngestPipeline,
    meta: SQLiteMeta,
    store: LanceDBStore,
    settings: Settings,
    collection: Collection,
    pdf_builder,
) -> None:
    doc = pipeline.ingest_bytes(collection.id, pdf_builder([[], [], []]), "scan.pdf").document
    assert doc.status == DocumentStatus.failed and doc.chunk_count == 0
    assert doc.error.startswith("This PDF has no extractable text")
    assert doc.title == "Scan"
    assert store.chunks_for(doc.id) == []
    assert (settings.uploads_dir / collection.id / doc.id / "scan.pdf").exists(), (
        "kept for download and retry"
    )
    assert meta.list_audit(event="ingest")[0]["outcome"] == "failed"


def test_retry_after_failure_reuses_the_document(
    flaky_pipeline: IngestPipeline, flaky: FlakyEmbedder, meta: SQLiteMeta, collection: Collection
) -> None:
    flaky.fail = True
    failed = flaky_pipeline.ingest_bytes(collection.id, V1, "travel.md").document
    assert failed.status == DocumentStatus.failed and failed.chunk_count == 0
    assert "previous version" not in failed.error
    flaky.fail = False
    retried = flaky_pipeline.ingest_bytes(collection.id, V1, "travel.md").document
    assert retried.id == failed.id and retried.status == DocumentStatus.ready
    assert len(meta.list_documents(collection.id)) == 1
    again = flaky_pipeline.ingest_bytes(collection.id, V1, "elsewhere/travel-copy.md")
    assert again.already_indexed and again.document.id == failed.id


def test_reindex_retries_a_failed_upload(
    flaky_pipeline: IngestPipeline, flaky: FlakyEmbedder, collection: Collection
) -> None:
    flaky.fail = True
    failed = flaky_pipeline.ingest_bytes(collection.id, V1, "travel.md").document
    flaky.fail = False
    retried = flaky_pipeline.reindex(failed.id).document
    assert retried.status == DocumentStatus.ready and retried.title == "Travel Rules", "the parsed title applies"


# ----------------------------------------------------------------------------------- failure rules


def test_refused_before_any_work_leaves_no_row(
    pipeline: IngestPipeline, meta: SQLiteMeta, settings: Settings, collection: Collection, pdf_builder
) -> None:
    settings.MAX_UPLOAD_MB = 1
    locked_pdf = encrypted_pdf(pdf_builder([["Secret text for the test."]]))
    locked_docx = bytes.fromhex("d0cf11e0a1b11ae1") + bytes(504) + "EncryptedPackage".encode("utf-16-le")
    with pytest.raises(UnsupportedFileTypeError):
        pipeline.ingest_bytes(collection.id, b"MZ", "tool.exe")
    with pytest.raises(EmptyDocumentError):
        pipeline.ingest_bytes(collection.id, b"", "empty.md")
    with pytest.raises(FileTooLargeError, match="1 MB"):
        pipeline.ingest_bytes(collection.id, b"word " * 250_000, "big.txt")
    with pytest.raises(EncryptedFileError):
        pipeline.ingest_bytes(collection.id, locked_pdf, "locked.pdf")
    with pytest.raises(EncryptedFileError):
        pipeline.ingest_bytes(collection.id, locked_docx, "locked.docx")
    with pytest.raises(NotFoundError):
        pipeline.ingest_bytes("no-such-collection", V1, "travel.md")
    assert meta.list_documents(collection.id) == []
    assert meta.current_version() == ""
    staging = settings.uploads_dir / collection.id / ".staging"
    assert not staging.exists() or not any(staging.iterdir())


def _invalid_docx() -> bytes:
    return b"PK\x03\x04 this is not a Word file"


@pytest.mark.parametrize(
    ("filename", "make", "message"),
    [
        ("blank.txt", lambda pdf: b"   \n\n \t \n", "no readable text"),
        ("broken.pdf", lambda pdf: b"%PDF-1.4 this is not really a pdf", "could not be opened"),
        ("broken.docx", lambda pdf: _invalid_docx(), "Word file could not be read"),
        ("scan.pdf", lambda pdf: pdf([[], [], []]), "scanned"),
    ],
)
def test_accepted_then_failing_leaves_a_failed_row(
    pipeline: IngestPipeline,
    settings: Settings,
    collection: Collection,
    pdf_builder,
    filename: str,
    make: Any,
    message: str,
) -> None:
    doc = pipeline.ingest_bytes(collection.id, make(pdf_builder), filename).document
    assert doc.status == DocumentStatus.failed and message in doc.error
    assert "/" not in doc.error and "\\" not in doc.error and "Traceback" not in doc.error
    assert pipeline.original_path(doc).exists(), "kept so the person can retry or dismiss it"


# ------------------------------------------------------------------------------- progress and batches


def test_progress_reports_every_stage_and_every_batch(
    gated_pipeline: IngestPipeline, gated: GatedEmbedder, settings: Settings, collection: Collection
) -> None:
    settings.EMBED_BATCH_SIZE = 4
    updates: list[dict[str, Any]] = []
    doc = gated_pipeline.ingest_bytes(collection.id, V1, "travel.md", progress=updates.append).document
    total = doc.chunk_count
    assert total > 4 and all(n <= 4 for n in gated.batches) and sum(gated.batches) == total
    stages = [u["stage"] for u in updates]
    assert stages[:3] == ["parsing", "chunking", "embedding"] and stages[-2:] == ["writing", "done"]
    embedding = [u for u in updates if u["stage"] == "embedding"]
    assert len(embedding) == 1 + len(gated.batches), "one at the start and one after every batch"
    assert [u["chunks_done"] for u in embedding] == sorted(u["chunks_done"] for u in embedding)
    assert embedding[-1]["chunks_done"] == total and all(u["chunks_total"] == total for u in embedding)
    assert all(isinstance(u["eta_seconds"], int) and u["eta_seconds"] >= 0 for u in embedding[1:])
    assert updates[-1] == {"stage": "done", "chunks_done": total, "chunks_total": total, "eta_seconds": 0}
    assert set(updates[0]) == {"stage", "chunks_done", "chunks_total", "eta_seconds"}


def test_a_failing_progress_callback_does_not_stop_indexing(
    pipeline: IngestPipeline, collection: Collection
) -> None:
    def broken(update: dict[str, Any]) -> None:
        raise RuntimeError("the UI went away")

    assert pipeline.ingest_bytes(collection.id, V1, "travel.md", progress=broken).document.status == "ready"


def test_progress_is_written_to_the_job(pipeline: IngestPipeline, meta: SQLiteMeta, collection: Collection) -> None:
    job = meta.create_job(collection.id, "", filename="travel.md", bytes_total=len(V1))
    seen: list[str] = []

    def to_job(update: dict[str, Any]) -> None:
        seen.append(meta.update_job(job.id, **update).stage)

    pipeline.ingest_bytes(collection.id, V1, "travel.md", job_id=job.id, progress=to_job)
    final = meta.get_job(job.id)
    assert seen[-1] == "done" and final.status == DocumentStatus.queued, "progress alone never sets a status"
    assert final.chunks_done == final.chunks_total and final.eta_seconds == 0


# ------------------------------------------------------------------------------------- cancellation


def test_cancel_mid_embedding_keeps_the_previous_version(
    pipeline: IngestPipeline,
    meta: SQLiteMeta,
    store: LanceDBStore,
    settings: Settings,
    collection: Collection,
) -> None:
    settings.EMBED_BATCH_SIZE = 2
    v1 = pipeline.ingest_bytes(collection.id, V1, "travel.md").document
    version = meta.current_version()
    job = meta.create_job(collection.id, v1.id)
    updates: list[dict[str, Any]] = []

    def cancel_after_first_batch(update: dict[str, Any]) -> None:
        updates.append(update)
        if update["stage"] == "embedding" and update["chunks_done"] > 0:
            meta.request_cancel(job.id)

    v2 = markdown("Travel Rules", {"Flights": "Business class for everyone. " + filler("flights", 30)})
    with pytest.raises(IngestCancelled):
        pipeline.ingest_bytes(collection.id, v2, "travel.md", job_id=job.id, progress=cancel_after_first_batch)
    assert [u["stage"] for u in updates][-1] == "embedding", "it stopped before writing"
    after = meta.get_document(v1.id)
    assert after == v1, "status, hash, title and counts are those of the previous version"
    assert len(store.chunks_for(v1.id)) == v1.chunk_count and "Business class" not in texts(store, v1.id)
    assert meta.current_version() == version
    assert (settings.uploads_dir / collection.id / v1.id / "travel.md").read_bytes() == V1
    assert meta.list_audit(event="ingest")[0]["outcome"] == "cancelled"
    assert pipeline.find_duplicate(collection.id, v2) is None, "nothing is left in flight"


def test_cancel_of_a_new_upload_leaves_no_document(
    pipeline: IngestPipeline, meta: SQLiteMeta, store: LanceDBStore, settings: Settings, collection: Collection
) -> None:
    job = meta.create_job(collection.id, "")
    meta.request_cancel(job.id)
    with pytest.raises(IngestCancelled):
        pipeline.ingest_bytes(collection.id, V1, "travel.md", job_id=job.id)
    assert meta.list_documents(collection.id) == [] and store.count(collection.id) == 0
    folder = settings.uploads_dir / collection.id
    leftovers = [p.name for p in folder.iterdir()] if folder.exists() else []
    assert leftovers in ([], [".staging"]), leftovers
    assert pipeline.ingest_bytes(collection.id, V1, "travel.md").document.status == DocumentStatus.ready


def test_cancel_of_a_reindex_keeps_the_index(
    pipeline: IngestPipeline, meta: SQLiteMeta, store: LanceDBStore, collection: Collection
) -> None:
    doc = pipeline.ingest_bytes(collection.id, V1, "travel.md").document
    job = meta.create_job(collection.id, doc.id)
    meta.request_cancel(job.id)
    with pytest.raises(IngestCancelled):
        pipeline.reindex(doc.id, job_id=job.id)
    assert meta.get_document(doc.id) == doc and len(store.chunks_for(doc.id)) == doc.chunk_count


# ------------------------------------------------------------------------------------- concurrency


def test_three_identical_uploads_at_once_make_one_document(
    gated_pipeline: IngestPipeline, gated: GatedEmbedder, meta: SQLiteMeta, collection: Collection
) -> None:
    gated.gate.clear()
    first, first_out = in_thread(gated_pipeline.ingest_bytes, collection.id, V1, "travel.md")
    assert gated.waiting.wait(10), "the first upload reached embedding"
    in_flight = gated_pipeline.find_duplicate(collection.id, V1)
    assert in_flight is not None and in_flight.status == DocumentStatus.processing
    others = [in_thread(gated_pipeline.ingest_bytes, collection.id, V1, name) for name in ("a.md", "b.md")]
    for thread, _ in others:
        thread.join(10)
    gated.gate.set()
    first.join(10)
    results = [first_out, *(out for _, out in others)]
    assert all("error" not in out for out in results), results
    assert [out["result"].already_indexed for out in results] == [False, True, True]
    assert len({out["result"].document.id for out in results}) == 1
    assert len(meta.list_documents(collection.id)) == 1
    assert meta.list_documents(collection.id)[0].status == DocumentStatus.ready


def test_delete_reindex_and_upload_never_wait_for_another_documents_embedding(
    gated_pipeline: IngestPipeline, gated: GatedEmbedder, meta: SQLiteMeta, store: LanceDBStore, collection: Collection
) -> None:
    """QA saw a delete wait two minutes behind one upload. Now only the store's short writes are shared."""
    other = gated_pipeline.ingest_bytes(collection.id, V2, "other.md").document
    third = gated_pipeline.ingest_bytes(collection.id, markdown("Third", {"Body": filler("desks", 4)}), "c.md").document
    gated.block_on = "zebrafruit"  # only the travel upload below waits
    gated.gate.clear()
    busy, busy_out = in_thread(gated_pipeline.ingest_bytes, collection.id, V1, "travel.md")
    assert gated.waiting.wait(10)

    for action, args in [
        (gated_pipeline.delete, (other.id,)),
        (gated_pipeline.reindex, (third.id,)),
        (gated_pipeline.ingest_bytes, (collection.id, V2, "other-again.md")),
    ]:
        thread, out = in_thread(action, *args)
        thread.join(5)
        assert not thread.is_alive(), f"{action.__name__} waited for another document's embedding"
        assert "error" not in out, out
    assert busy.is_alive(), "the travel upload was still embedding all along"
    assert meta.get_document(other.id) is None and store.chunks_for(other.id) == []
    assert meta.get_document(third.id).status == DocumentStatus.ready
    gated.gate.set()
    busy.join(10)
    assert busy_out["result"].document.status == DocumentStatus.ready
    assert len(meta.list_documents(collection.id)) == 3


def test_deleting_a_document_while_it_is_indexed_stops_the_upload(
    gated_pipeline: IngestPipeline, gated: GatedEmbedder, meta: SQLiteMeta, store: LanceDBStore, collection: Collection
) -> None:
    gated.gate.clear()
    busy, busy_out = in_thread(gated_pipeline.ingest_bytes, collection.id, V1, "travel.md")
    assert gated.waiting.wait(10)
    [doc] = meta.list_documents(collection.id)
    deleting, delete_out = in_thread(gated_pipeline.delete, doc.id)
    time.sleep(0.2)
    gated.gate.set()
    deleting.join(10)
    busy.join(10)
    assert "error" not in delete_out
    assert isinstance(busy_out.get("error"), IngestCancelled)
    assert meta.list_documents(collection.id) == [] and store.count(collection.id) == 0


def test_deleting_a_collection_stops_its_uploads(
    gated_pipeline: IngestPipeline, gated: GatedEmbedder, meta: SQLiteMeta, collection: Collection
) -> None:
    gated.gate.clear()
    busy, busy_out = in_thread(gated_pipeline.ingest_bytes, collection.id, V1, "travel.md")
    assert gated.waiting.wait(10)
    deleting, delete_out = in_thread(gated_pipeline.delete_collection, collection.id)
    time.sleep(0.2)
    with pytest.raises(NotFoundError):
        gated_pipeline.ingest_bytes(collection.id, V2, "late.md")
    gated.gate.set()
    deleting.join(10)
    busy.join(10)
    assert "error" not in delete_out and isinstance(busy_out.get("error"), IngestCancelled)
    assert meta.get_collection(collection.id) is None and meta.list_documents(collection.id) == []


def test_concurrent_ingestion_of_different_files_is_consistent(
    pipeline: IngestPipeline, meta: SQLiteMeta, collection: Collection
) -> None:
    results: dict[int, IngestResult] = {}
    errors: list[BaseException] = []

    def ingest(n: int) -> None:
        try:
            data = markdown(f"Doc {n}", {"Body": f"Unique body {n}. " + filler(f"topic {n}", 3)})
            results[n] = pipeline.ingest_bytes(collection.id, data, f"doc{n}.md")
        except BaseException as exc:  # noqa: BLE001 - report any failure from the thread
            errors.append(exc)

    threads = [threading.Thread(target=ingest, args=(n,)) for n in range(4)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert errors == []
    assert all(r.document.status == DocumentStatus.ready for r in results.values())
    assert len({r.document.knowledge_version for r in results.values()}) == 4
    assert len(meta.list_documents(collection.id)) == 4


# ------------------------------------------------------------------------------- pages, sheets, size


def test_every_fact_of_a_40_page_pdf_carries_its_page(
    pipeline: IngestPipeline, store: LanceDBStore, collection: Collection, pdf_builder
) -> None:
    doc = pipeline.ingest_bytes(collection.id, fact_pdf(pdf_builder, 40), "vaults.pdf").document
    assert doc.status == DocumentStatus.ready and doc.pages == 40
    chunks = store.chunks_for(doc.id)
    for page, word in enumerate(CODE_WORDS, start=1):
        [chunk] = [c for c in chunks if word in c.text]
        assert (chunk.page, chunk.page_start, chunk.page_end) == (page, page, page), word
        [hit] = [h for h in store.text_search(collection.id, word, 3) if word in h.chunk.text]
        assert hit.chunk.page == page


def test_a_three_page_pdf_is_three_chunks(
    pipeline: IngestPipeline, store: LanceDBStore, collection: Collection, pdf_builder
) -> None:
    doc = pipeline.ingest_bytes(collection.id, fact_pdf(pdf_builder, 3), "three.pdf").document
    assert [c.page for c in store.chunks_for(doc.id)] == [1, 2, 3]


def test_hidden_sheets_are_not_answerable(
    pipeline: IngestPipeline, store: LanceDBStore, settings: Settings, collection: Collection, tmp_path: Path
) -> None:
    wb = Workbook()
    wb.active.title = "Prices"
    wb.active.append(["Item", "Price"])
    wb.active.append(["Desk", 120])
    secret = wb.create_sheet("Salaries")
    secret.append(["Name", "Salary"])
    secret.append(["Quentin Quarrington", 250000])
    secret.sheet_state = "hidden"
    wb.save(str(tmp_path / "book.xlsx"))
    data = (tmp_path / "book.xlsx").read_bytes()
    doc = pipeline.ingest_bytes(collection.id, data, "book.xlsx").document
    assert doc.status == DocumentStatus.ready
    assert store.text_search(collection.id, "Quarrington", 5) == []
    assert store.text_search(collection.id, "Desk", 5)[0].chunk.section == "Prices"


def test_a_200_kb_line_is_cut_to_the_character_ceiling(
    pipeline: IngestPipeline, store: LanceDBStore, settings: Settings, collection: Collection
) -> None:
    data = ("Header line\n\n" + "z" * 200_000 + "\n").encode()
    doc = pipeline.ingest_bytes(collection.id, data, "blob.txt").document
    chunks = store.chunks_for(doc.id)
    assert doc.status == DocumentStatus.ready and len(chunks) >= 50
    assert all(len(c.text) <= settings.MAX_CHUNK_CHARS for c in chunks)


def test_word_document_ingests(pipeline: IngestPipeline, collection: Collection, tmp_path: Path) -> None:
    document = WordDocument()
    document.add_paragraph("Expense claims are filed within thirty days of travel.")
    document.save(str(tmp_path / "claims.docx"))
    doc = pipeline.ingest_bytes(collection.id, (tmp_path / "claims.docx").read_bytes(), "claims.docx").document
    assert doc.status == DocumentStatus.ready and doc.chunk_count == 1


def test_ingest_file_copies_the_original(
    pipeline: IngestPipeline, collection: Collection, sample_corpus_dir: Path
) -> None:
    source = sample_corpus_dir / "employee-handbook.md"
    before = source.read_bytes()
    doc = pipeline.ingest_file(collection.id, source).document
    assert doc.status == DocumentStatus.ready and doc.title == "Employee Handbook"
    assert doc.source_path == "employee-handbook.md"
    assert pipeline.original_path(doc).read_bytes() == before
    assert source.read_bytes() == before


# ------------------------------------------------------------------------------ reindex and delete


def test_reindex_from_the_kept_original(
    pipeline: IngestPipeline, meta: SQLiteMeta, store: LanceDBStore, collection: Collection
) -> None:
    doc = pipeline.ingest_bytes(collection.id, V1, "travel.md").document
    ids = [c.id for c in store.chunks_for(doc.id)]
    result = pipeline.reindex(doc.id)
    again = result.document
    assert (result.already_indexed, result.replaced_doc_id) == (False, None)
    assert again.id == doc.id and again.status == DocumentStatus.ready
    assert again.knowledge_version != doc.knowledge_version
    assert [c.id for c in store.chunks_for(doc.id)] == ids
    assert {c.knowledge_version for c in store.chunks_for(doc.id)} == {again.knowledge_version}
    assert meta.list_audit(event="reindex")[0]["outcome"] == "indexed"
    with pytest.raises(NotFoundError):
        pipeline.reindex("missing")


def test_reindex_keeps_an_explicit_title(pipeline: IngestPipeline, collection: Collection) -> None:
    doc = pipeline.ingest_bytes(collection.id, V1, "travel.md", title="My travel notes").document
    assert pipeline.reindex(doc.id).document.title == "My travel notes"


def test_reindex_without_the_original_keeps_the_index(
    pipeline: IngestPipeline, store: LanceDBStore, collection: Collection
) -> None:
    doc = pipeline.ingest_bytes(collection.id, V1, "travel.md").document
    pipeline.original_path(doc).unlink()
    failed = pipeline.reindex(doc.id).document
    assert failed.status == DocumentStatus.failed and "original file is missing" in failed.error
    assert "previous index is still searchable" in failed.error
    assert len(store.chunks_for(doc.id)) == doc.chunk_count


def test_delete_removes_chunks_metadata_and_original(
    pipeline: IngestPipeline,
    meta: SQLiteMeta,
    store: LanceDBStore,
    settings: Settings,
    collection: Collection,
) -> None:
    keep = pipeline.ingest_bytes(collection.id, V2, "keep.md", title="Keep me").document
    gone = pipeline.ingest_bytes(collection.id, V1, "travel.md").document
    version = meta.current_version()
    pipeline.delete(gone.id)
    assert store.chunks_for(gone.id) == [] and meta.get_document(gone.id) is None
    assert not (settings.uploads_dir / collection.id / gone.id).exists()
    assert len(store.chunks_for(keep.id)) == keep.chunk_count and meta.get_document(keep.id) is not None
    assert meta.current_version() != version
    audit = meta.list_audit(event="delete")[0]
    assert audit["doc_id"] == gone.id and audit["chunks"] == gone.chunk_count
    with pytest.raises(NotFoundError):
        pipeline.delete(gone.id)


def test_delete_collection_removes_only_that_collection(
    pipeline: IngestPipeline,
    meta: SQLiteMeta,
    store: LanceDBStore,
    settings: Settings,
    collection: Collection,
) -> None:
    other = meta.create_collection("Other")
    pipeline.ingest_bytes(collection.id, V1, "travel.md")
    survivor = pipeline.ingest_bytes(other.id, V2, "travel.md").document
    removed = pipeline.delete_collection(collection.id)
    assert removed > 0 and store.count(collection.id) == 0
    assert meta.get_collection(collection.id) is None and meta.list_documents(collection.id) == []
    assert not (settings.uploads_dir / collection.id).exists()
    assert store.count(other.id) == survivor.chunk_count
    with pytest.raises(NotFoundError):
        pipeline.delete_collection(collection.id)


# ------------------------------------------------------------------------------- versions and names


def test_every_change_bumps_the_version(
    pipeline: IngestPipeline, meta: SQLiteMeta, collection: Collection
) -> None:
    a = pipeline.ingest_bytes(collection.id, V1, "a.md", title="A").document
    b = pipeline.ingest_bytes(collection.id, V2, "b.md", title="B").document
    day, n = a.knowledge_version.split(".")
    assert VERSION.fullmatch(a.knowledge_version) and b.knowledge_version == f"{day}.{int(n) + 1}"
    pipeline.delete(a.id)
    assert meta.current_version() == f"{day}.{int(n) + 2}"


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("report.pdf", "report.pdf"),
        ("../../etc/passwd.md", "passwd.md"),
        ("C:\\Users\\me\\notes.txt", "notes.txt"),
        ('bad<>:"|?*name.md', "bad_______name.md"),
        ("CON.txt", "_CON.txt"),
        (".md", "upload.md"),
        ("trailing. ", "trailing"),
        ("", "upload"),
    ],
)
def test_safe_filename(raw: str, expected: str) -> None:
    assert safe_filename(raw) == expected


def test_path_traversal_upload_stays_inside_the_uploads_folder(
    pipeline: IngestPipeline, settings: Settings, collection: Collection
) -> None:
    doc = pipeline.ingest_bytes(collection.id, V1, "../../../outside.md").document
    path = pipeline.original_path(doc).resolve()
    assert settings.uploads_dir.resolve() in path.parents and path.name == "outside.md"


def test_the_whole_sample_corpus_indexes(
    indexed_corpus: Collection, meta: SQLiteMeta, store: LanceDBStore
) -> None:
    docs = meta.list_documents(indexed_corpus.id)
    assert len(docs) >= 8
    assert all(d.status == DocumentStatus.ready and d.chunk_count > 0 for d in docs)
    assert store.count(indexed_corpus.id) == sum(d.chunk_count for d in docs)
    pdf = next(d for d in docs if d.media_type == "application/pdf")
    assert pdf.pages == 2
