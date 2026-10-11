"""Parsers: one file in, structural blocks out, for every supported format."""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path

import pytest
from docx import Document as WordDocument
from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter

from ragkb.ingest.parsers import (
    SUPPORTED_SUFFIXES,
    EmptyDocumentError,
    EncryptedFileError,
    ParseError,
    UnsupportedFileTypeError,
    check_not_encrypted,
    media_type_for,
    parse,
    title_from_filename,
)

BODY = [
    "The regional teams reported fewer incidents this year than in any year before, and the board",
    "asked every site to keep the weekly safety walk that made the difference in the first quarter.",
    "Contractors follow the same induction as staff, which now takes two hours instead of four.",
]


def encrypted_pdf(data: bytes, user_password: str = "secret") -> bytes:
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(data)))
    writer.encrypt(user_password, owner_password="owner", algorithm="RC4-128")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def write(path: Path, text: str) -> Path:
    path.write_bytes(text.encode("utf-8"))
    return path


def _write_bytes(path: Path, data: bytes) -> Path:
    path.write_bytes(data)
    return path


# ------------------------------------------------------------------------------------- media types


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        ("a.pdf", "application/pdf"),
        ("a.DOCX", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ("a.md", "text/markdown"),
        ("a.markdown", "text/markdown"),
        ("a.html", "text/html"),
        ("a.htm", "text/html"),
        ("a.txt", "text/plain"),
        ("a.csv", "text/csv"),
        ("a.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ("a.exe", None),
        ("a.doc", None),
        ("noextension", None),
    ],
)
def test_media_type_for(name: str, expected: str | None) -> None:
    assert media_type_for(name) == expected


def test_supported_suffixes_cover_the_promised_formats() -> None:
    assert {".pdf", ".docx", ".md", ".txt", ".html", ".csv", ".xlsx"} <= SUPPORTED_SUFFIXES


def test_unsupported_type_raises_plain_message(tmp_path: Path) -> None:
    with pytest.raises(UnsupportedFileTypeError) as err:
        parse(write(tmp_path / "x.exe", "MZ"))
    assert ".exe" in err.value.message and ".pdf" in err.value.message
    assert err.value.code == "unsupported_file_type"


def test_legacy_word_file_gets_a_hint(tmp_path: Path) -> None:
    with pytest.raises(UnsupportedFileTypeError, match=r"\.docx"):
        parse(write(tmp_path / "old.doc", "binary"))


def test_title_from_filename() -> None:
    assert title_from_filename("travel-and_expense-policy.md") == "Travel and expense policy"
    assert title_from_filename("Q3 Report.pdf") == "Q3 Report"
    assert title_from_filename("___.txt") == "Untitled"


# ---------------------------------------------------------------------------------------- markdown


MARKDOWN = """---
author: someone
---
# Expense Guide

Read this **before** you travel. See [the portal](https://example.invalid) for forms.

## Hotels

### Caps

Tier 1 cities are capped at EUR 210 per night.

- Book through the desk
- Keep receipts
  - including taxis
1. First step
2. Second step

| Tier | Cap |
|---|---|
| 1 | 210 |
| 2 | 150 |

```
claim --submit
```

## Meals

Per diem is EUR 55.
"""


def test_markdown_structure(tmp_path: Path) -> None:
    doc = parse(write(tmp_path / "guide.md", MARKDOWN))
    assert doc.title == "Expense Guide"
    assert doc.media_type == "text/markdown"
    assert doc.pages is None
    kinds = [(b.kind, b.heading_path) for b in doc.blocks]
    assert kinds[0] == ("paragraph", [])
    assert ("heading", ["Hotels"]) in kinds
    assert ("heading", ["Hotels", "Caps"]) in kinds
    assert ("heading", ["Meals"]) in kinds
    intro = doc.blocks[0].text
    assert intro == "Read this before you travel. See the portal for forms."
    lists = [b for b in doc.blocks if b.kind == "list"]
    assert lists[0].text.splitlines() == ["- Book through the desk", "- Keep receipts", "  - including taxis"]
    assert lists[1].text.splitlines() == ["1. First step", "2. Second step"]
    table = next(b for b in doc.blocks if b.kind == "table")
    assert table.text.splitlines() == ["Tier | Cap", "1 | 210", "2 | 150"]
    assert table.heading_path == ["Hotels", "Caps"]
    code = next(b for b in doc.blocks if b.kind == "code")
    assert code.text == "claim --submit"
    meals = doc.blocks[-1]
    assert meals.text == "Per diem is EUR 55." and meals.heading_path == ["Meals"]


def test_markdown_front_matter_title_wins(tmp_path: Path) -> None:
    doc = parse(
        write(tmp_path / "x.md", "---\ntitle: Official Name\n---\n# Heading One\n\nBody text here.\n")
    )
    assert doc.title == "Official Name"
    assert all(b.text != "Heading One" for b in doc.blocks)


def test_markdown_without_title_uses_filename(tmp_path: Path) -> None:
    doc = parse(write(tmp_path / "remote-work.md", "Body first.\n\n# Later heading\n\nMore.\n"))
    assert doc.title == "Remote work"
    assert doc.blocks[1].kind == "heading" and doc.blocks[1].text == "Later heading"


# -------------------------------------------------------------------------------------------- html


HTML = """<html><head><title>Guide | Intranet</title><style>p { color: red }</style>
<script>var secret = "do not index";</script></head>
<body><nav><a href="/">Home</a> Menu items</nav>
<main>
<h1>Onboarding</h1>
<p>Welcome<br>to the team.</p>
<h2>First day</h2>
<div>Arrive at <b>09:00</b> sharp.<p>Bring ID.</p></div>
<ul><li>Badge</li><li>Laptop<ul><li>Charger</li></ul></li></ul>
<ol><li>One</li><li>Two</li></ol>
<table><tr><th>Day</th><th>Who</th></tr><tr><td>30</td><td>Manager</td></tr></table>
<pre>code  stays</pre>
</main>
<footer>Copyright footer</footer>
</body></html>"""


def test_html_structure(tmp_path: Path) -> None:
    doc = parse(write(tmp_path / "page.html", HTML))
    assert doc.title == "Onboarding"
    texts = [b.text for b in doc.blocks]
    joined = "\n".join(texts)
    for removed in ("do not index", "Menu items", "Copyright footer", "color: red"):
        assert removed not in joined
    assert texts[0] == "Welcome\nto the team."
    assert doc.blocks[0].heading_path == []
    first_day = [b for b in doc.blocks if b.heading_path == ["First day"]]
    assert [b.kind for b in first_day] == [
        "heading",
        "paragraph",
        "paragraph",
        "list",
        "list",
        "table",
        "code",
    ]
    assert first_day[1].text == "Arrive at 09:00 sharp."
    assert first_day[3].text.splitlines() == ["- Badge", "- Laptop", "  - Charger"]
    assert first_day[4].text.splitlines() == ["1. One", "2. Two"]
    assert first_day[5].text == "Day | Who\n30 | Manager"
    assert first_day[6].text == "code  stays"


def test_html_title_falls_back_to_title_tag(tmp_path: Path) -> None:
    doc = parse(
        write(
            tmp_path / "p.html",
            "<html><head><title>Page Title</title></head><body><p>Hi there.</p></body></html>",
        )
    )
    assert doc.title == "Page Title"


# -------------------------------------------------------------------------------------------- docx


def make_docx(path: Path) -> Path:
    document = WordDocument()
    document.core_properties.title = "Properties Title"
    document.add_paragraph("Retention Policy", style="Title")
    document.add_paragraph("Intro paragraph about retention.")
    document.add_heading("Periods", level=1)
    table = document.add_table(rows=3, cols=2)
    for r, row in enumerate([("Record", "Keep for"), ("Invoices", "7 years"), ("CCTV", "30 days")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    document.add_heading("Hold", level=2)
    document.add_paragraph("First rule", style="List Bullet")
    document.add_paragraph("Second rule", style="List Bullet")
    document.add_paragraph("Closing words.")
    document.save(str(path))
    return path


def test_docx_structure(tmp_path: Path) -> None:
    doc = parse(make_docx(tmp_path / "policy.docx"))
    assert doc.title == "Retention Policy"
    assert [(b.kind, b.heading_path) for b in doc.blocks] == [
        ("paragraph", []),
        ("heading", ["Periods"]),
        ("table", ["Periods"]),
        ("heading", ["Periods", "Hold"]),
        ("list", ["Periods", "Hold"]),
        ("paragraph", ["Periods", "Hold"]),
    ]
    assert doc.blocks[2].text == "Record | Keep for\nInvoices | 7 years\nCCTV | 30 days"
    assert doc.blocks[4].text == "- First rule\n- Second rule"


def test_docx_merged_cells_are_not_repeated(tmp_path: Path) -> None:
    document = WordDocument()
    table = document.add_table(rows=2, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "Wide header"
    table.cell(0, 2).text = "Other"
    table.cell(1, 0).text = "a"
    table.cell(1, 1).text = "b"
    table.cell(1, 2).text = "c"
    document.save(str(tmp_path / "m.docx"))
    doc = parse(tmp_path / "m.docx")
    assert doc.blocks[0].text.splitlines() == ["Wide header | Other", "a | b | c"]
    assert doc.title == "M"


def test_damaged_docx_raises_parse_error(tmp_path: Path) -> None:
    with pytest.raises(ParseError, match="Word file could not be read"):
        parse(write(tmp_path / "broken.docx", "this is not a zip file"))


# --------------------------------------------------------------------------------------------- pdf


def test_pdf_pages_and_paragraphs(tmp_path: Path, pdf_builder) -> None:
    data = pdf_builder(
        [
            [
                "Benefits Overview",
                "",
                "The company matches pension contributions up to 6 percent of base salary for every",
                "permanent employee. Contributions can be changed in the portal at any time, and the",
                "change applies from the next payroll run.",
                "",
                "Health insurance",
                "",
                "Dental cover is included up to EUR 600 per person per year. Partner cover costs",
                "EUR 35 per month. The insurer handles claims online with a turnaround of five days.",
            ],
            [
                "Learning budget",
                "",
                "Every employee has EUR 1,200 per year for courses, books and certification fees. The",
                "budget does not carry over, and requests above EUR 500 need manager approval through",
                "the learning portal, which also tracks the remaining budget for each employee.",
            ],
        ],
        title="Benefits Overview",
    )
    doc = parse(_write_bytes(tmp_path / "benefits.pdf", data))
    assert doc.pages == 2
    assert doc.title == "Benefits Overview"
    paragraphs = [b for b in doc.blocks if b.kind == "paragraph"]
    assert paragraphs[0].page == 1 and paragraphs[0].text.startswith("The company matches pension")
    assert paragraphs[0].text.endswith("from the next payroll run.")
    assert paragraphs[-1].page == 2 and "EUR 1,200" in paragraphs[-1].text
    headings = [b for b in doc.blocks if b.kind == "heading"]
    assert [(h.text, h.page) for h in headings] == [("Health insurance", 1), ("Learning budget", 2)]
    assert paragraphs[1].heading_path == ["Health insurance"]


def test_pdf_joins_hyphenated_line_breaks_and_drops_page_numbers(tmp_path: Path, pdf_builder) -> None:
    data = pdf_builder(
        [
            [
                "Records are reviewed by the information manage-",
                "ment team every quarter, and the outcome of the review is reported to the board.",
                "Page 1 of 1",
            ]
        ]
    )
    doc = parse(_write_bytes(tmp_path / "h.pdf", data))
    text = " ".join(b.text for b in doc.blocks)
    assert "information management team" in text
    assert "Page 1 of 1" not in text


def test_scanned_pdf_is_rejected_as_empty(tmp_path: Path, pdf_builder) -> None:
    data = pdf_builder([[], [], []])
    with pytest.raises(EmptyDocumentError, match="scanned"):
        parse(_write_bytes(tmp_path / "scan.pdf", data))


def test_pdf_with_almost_no_text_is_rejected(tmp_path: Path, pdf_builder) -> None:
    data = pdf_builder([["x"], [], [], [], []])
    with pytest.raises(EmptyDocumentError):
        parse(_write_bytes(tmp_path / "scan.pdf", data))


def test_damaged_pdf_raises_parse_error(tmp_path: Path) -> None:
    path = _write_bytes(tmp_path / "bad.pdf", b"%PDF-1.4 this is not really a pdf")
    check_not_encrypted(path)  # a damaged file is not refused up front; it fails when parsed
    with pytest.raises(ParseError):
        parse(path)


def test_pdf_title_is_the_first_line_when_it_reads_as_a_title(tmp_path: Path, pdf_builder) -> None:
    long_title = "Annual Safety Performance Report for the Northern Region Operations 2026"
    doc = parse(_write_bytes(tmp_path / "report-final-v3.pdf", pdf_builder([[long_title, "", *BODY]])))
    assert doc.title == long_title
    assert all(long_title not in b.text for b in doc.blocks), "the title leaves the text"
    # a title-case line directly followed by body text, with no blank line between them
    doc = parse(_write_bytes(tmp_path / "r.pdf", pdf_builder([["Quarterly Incident Review", *BODY]])))
    assert doc.title == "Quarterly Incident Review"


def test_pdf_first_line_that_is_a_sentence_is_not_a_title(tmp_path: Path, pdf_builder) -> None:
    doc = parse(_write_bytes(tmp_path / "site-notes.pdf", pdf_builder([[*BODY, "", *BODY]])))
    assert doc.title == "Site notes"
    doc = parse(_write_bytes(tmp_path / "short.pdf", pdf_builder([["All sites passed the audit.", "", *BODY]])))
    assert doc.title == "Short"


def test_pdf_metadata_title_wins(tmp_path: Path, pdf_builder) -> None:
    doc = parse(_write_bytes(tmp_path / "x.pdf", pdf_builder([["Some Heading Line", "", *BODY]], title="Official")))
    assert doc.title == "Official"
    assert doc.blocks[0].text == "Some Heading Line", "a different first line stays in the text"


def test_password_protected_pdf_is_refused(tmp_path: Path, pdf_builder) -> None:
    path = _write_bytes(tmp_path / "locked.pdf", encrypted_pdf(pdf_builder([["Secret text for the test.", *BODY]])))
    with pytest.raises(EncryptedFileError, match="password protected") as err:
        check_not_encrypted(path)
    assert err.value.code == "encrypted_file"
    with pytest.raises(EncryptedFileError):
        parse(path)
    opens = _write_bytes(tmp_path / "open.pdf", encrypted_pdf(pdf_builder([["Readable text for the test.", *BODY]]), ""))
    check_not_encrypted(opens)  # an empty user password opens it, so it is not refused
    assert "Readable text" in parse(opens).blocks[0].text


@pytest.mark.parametrize(("name", "label"), [("locked.docx", "Word"), ("locked.xlsx", "Excel")])
def test_password_protected_office_file_is_refused(tmp_path: Path, name: str, label: str) -> None:
    # An encrypted Office file is an OLE compound file with an "EncryptedPackage" stream.
    data = bytes.fromhex("d0cf11e0a1b11ae1") + bytes(504) + "EncryptedPackage".encode("utf-16-le") + bytes(64)
    with pytest.raises(EncryptedFileError, match=f"The {label} file is password protected"):
        check_not_encrypted(_write_bytes(tmp_path / name, data))
    check_not_encrypted(_write_bytes(tmp_path / "plain.docx", b"PK\x03\x04 not really a zip"))


# -------------------------------------------------------------------------------------- plain text


def test_text_title_lists_and_wrapped_paragraphs(tmp_path: Path) -> None:
    text = (
        "Security Rules\n\n"
        "Passwords must be long\nand unique for every account.\n\n"
        "Classes\n\n"
        "There are three classes:\n- Public: anyone\n- Internal: staff only\n  and contractors\n- Restricted: named people\n"
    )
    doc = parse(write(tmp_path / "rules.txt", text))
    assert doc.title == "Security Rules"
    assert doc.blocks[0].text == "Passwords must be long and unique for every account."
    assert doc.blocks[1].kind == "heading" and doc.blocks[1].text == "Classes"
    assert doc.blocks[2].text == "There are three classes:" and doc.blocks[2].heading_path == ["Classes"]
    assert doc.blocks[3].kind == "list"
    assert doc.blocks[3].text.splitlines() == [
        "- Public: anyone",
        "- Internal: staff only and contractors",
        "- Restricted: named people",
    ]


def test_text_in_cp1252_is_decoded(tmp_path: Path) -> None:
    (tmp_path / "legacy.txt").write_bytes("Caf\xe9 opening hours are 08:00 to 17:00.".encode("cp1252"))
    doc = parse(tmp_path / "legacy.txt")
    assert "Caf\xe9" in doc.blocks[0].text


@pytest.mark.parametrize("content", ["", "   \n\n  ", "--- ... ---"])
def test_empty_text_is_rejected(tmp_path: Path, content: str) -> None:
    with pytest.raises(EmptyDocumentError):
        parse(write(tmp_path / "empty.txt", content))


def test_empty_markdown_is_rejected(tmp_path: Path) -> None:
    with pytest.raises(EmptyDocumentError):
        parse(write(tmp_path / "empty.md", "# Only a title\n"))


# ------------------------------------------------------------------------------------ csv and xlsx


def test_csv_row_groups_repeat_the_header(tmp_path: Path) -> None:
    rows = ["Code,Description,Price"] + [
        f"SKU-{i},Replacement part number {i} for the sensor housing,{i * 3}" for i in range(80)
    ]
    doc = parse(write(tmp_path / "parts.csv", "\n".join(rows) + "\n"))
    assert len(doc.blocks) > 1
    assert all(b.kind == "table" for b in doc.blocks)
    for block in doc.blocks:
        assert block.text.splitlines()[0] == "Code | Description | Price"
    body = [line for b in doc.blocks for line in b.text.splitlines()[1:]]
    assert body == [
        f"SKU-{i} | Replacement part number {i} for the sensor housing | {i * 3}" for i in range(80)
    ]


def test_csv_semicolon_delimiter_is_detected(tmp_path: Path) -> None:
    doc = parse(write(tmp_path / "eu.csv", "Name;Amount\nTaxi;12,50\nHotel;140,00\n"))
    assert doc.blocks[0].text == "Name | Amount\nTaxi | 12,50\nHotel | 140,00"


def test_xlsx_sheets_and_values(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Prices"
    ws.append(["SKU", "Price", "Since", None])
    ws.append(["FH-1", 349.0, date(2026, 1, 15), None])
    ws.append([None, None, None, None])
    ws.append(["FH-2", 19.5, None, None])
    other = wb.create_sheet("Discounts")
    other.append(["Units", "Discount"])
    other.append(["10 to 49", "5%"])
    wb.save(str(tmp_path / "prices.xlsx"))
    doc = parse(tmp_path / "prices.xlsx")
    assert [(b.kind, b.text.splitlines()[0], b.heading_path) for b in doc.blocks] == [
        ("heading", "Prices", ["Prices"]),
        ("table", "SKU | Price | Since", ["Prices"]),
        ("heading", "Discounts", ["Discounts"]),
        ("table", "Units | Discount", ["Discounts"]),
    ]
    assert doc.blocks[1].text == "SKU | Price | Since\nFH-1 | 349 | 2026-01-15\nFH-2 | 19.5"
    assert doc.blocks[3].text == "Units | Discount\n10 to 49 | 5%"


def test_xlsx_hidden_sheets_are_left_out(tmp_path: Path) -> None:
    wb = Workbook()
    wb.active.title = "Public"
    wb.active.append(["Item", "Price"])
    wb.active.append(["Desk", 120])
    for name, state in (("Salaries", "hidden"), ("Margins", "veryHidden")):
        sheet = wb.create_sheet(name)
        sheet.append(["Name", "Amount"])
        sheet.append([f"secret {name}", 99000])
        sheet.sheet_state = state
    wb.save(str(tmp_path / "book.xlsx"))
    doc = parse(tmp_path / "book.xlsx")
    text = "\n".join(b.text for b in doc.blocks)
    assert "Desk" in text and "secret" not in text
    assert {tuple(b.heading_path) for b in doc.blocks} == {("Public",)}, "a named sheet keeps its heading"
    everything = parse(tmp_path / "book.xlsx", skip_hidden_sheets=False)
    assert "secret Salaries" in "\n".join(b.text for b in everything.blocks)


def test_xlsx_with_only_hidden_data_is_empty(tmp_path: Path) -> None:
    wb = Workbook()
    hidden = wb.create_sheet("Hidden data")
    hidden.append(["a", "b"])
    hidden.sheet_state = "hidden"
    wb.save(str(tmp_path / "hidden.xlsx"))
    with pytest.raises(EmptyDocumentError, match="no visible sheet"):
        parse(tmp_path / "hidden.xlsx")


def test_xlsx_single_default_sheet_has_no_path(tmp_path: Path) -> None:
    wb = Workbook()
    wb.active.append(["A", "B"])
    wb.active.append([1, 2])
    wb.save(str(tmp_path / "one.xlsx"))
    doc = parse(tmp_path / "one.xlsx")
    assert doc.blocks[0].heading_path == []


# ---------------------------------------------------------------------------------- sample corpus


def test_every_sample_document_parses(sample_corpus_dir: Path) -> None:
    files = sorted(p for p in sample_corpus_dir.iterdir() if p.suffix.lower() in SUPPORTED_SUFFIXES)
    assert len(files) >= 8
    suffixes = {p.suffix.lower() for p in files}
    assert {".md", ".txt", ".html", ".csv", ".docx", ".pdf", ".xlsx"} <= suffixes
    for path in files:
        doc = parse(path)
        assert doc.title and doc.blocks, path.name
        assert all(b.text.strip() for b in doc.blocks), path.name
